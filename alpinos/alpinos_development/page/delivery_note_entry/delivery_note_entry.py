import json

import frappe
from frappe.utils import cint, formatdate, getdate


@frappe.whitelist()
def get_delivery_note_data(name):
	"""Return Delivery Note header + items for the custom page."""
	dn = frappe.get_doc("Delivery Note", name)
	dn.check_permission("read")

	# Pick List name from the first item that carries one
	pick_list_name = ""
	for item in dn.items:
		if item.get("against_pick_list"):
			pick_list_name = item.against_pick_list
			break

	dispatch_date = ""
	if dn.custom_dispatch_date:
		try:
			dispatch_date = formatdate(str(dn.custom_dispatch_date)[:10])
		except Exception:
			dispatch_date = str(dn.custom_dispatch_date)

	dispatch_to_rows = []
	for row in (dn.get("custom_dispatch_to") or []):
		dispatch_to_rows.append({
			"name": row.name,
			"dispatch_to_address": row.get("dispatch_to_address") or "",
		})

	items = []
	for item in dn.items:
		items.append({
			"name": item.name,
			"item_code": item.item_code,
			"item_name": item.item_name,
			"qty": item.qty,
			"custom_box": item.get("custom_box") or 0,
			"batch_no": item.get("batch_no") or "",
			"custom_batch_code": item.get("custom_batch_code") or "",
			"custom_remark": item.get("custom_remark") or "",
			"custom_mfg_date": str(item.get("custom_mfg_date") or ""),
			"custom_expiry_date": str(item.get("custom_expiry_date") or ""),
			"against_pick_list": item.get("against_pick_list") or "",
		})

	return {
		"name": dn.name,
		"docstatus": dn.docstatus,
		"owner": dn.owner,
		"owner_full_name": frappe.utils.get_fullname(dn.owner),
		"posting_date": formatdate(str(dn.posting_date)) if dn.posting_date else "",
		"custom_sales_order_id": dn.get("custom_sales_order_id") or "",
		# Invoice No is set on the Sales Order; show its live value.
		"custom_invoice_no": (
			frappe.db.get_value("Sales Order", dn.get("custom_sales_order_id"), "custom_invoice_no")
			if dn.get("custom_sales_order_id") else ""
		) or "",
		"pick_list_name": pick_list_name,
		"custom_lr_gr_no": dn.get("custom_lr_gr_no") or "",
		"custom_dispatch_from": dn.get("custom_dispatch_from") or "",
		"custom_dn_so_customer_name": dn.get("custom_dn_so_customer_name") or "",
		"custom_transporter_name": dn.get("custom_transporter_name") or "",
		"vehicle_no": dn.get("vehicle_no") or "",
		"custom_dispatch_date": dispatch_date,
		# yyyy-mm-dd for the page's date input (Changes(HP) #43).
		"custom_dispatch_date_value": str(dn.custom_dispatch_date)[:10] if dn.custom_dispatch_date else "",
		"custom_assigned_to": dn.get("custom_assigned_to") or "",
		"custom_total_boxes": dn.get("custom_total_boxes") or 0,
		"custom_dn_order_gross_weight": dn.get("custom_dn_order_gross_weight") or 0,
		"custom_total_units_dn": dn.get("custom_total_units_dn") or 0,
		"items": items,
		"custom_dispatch_to": dispatch_to_rows,
	}


_EDITABLE_HEADER_FIELDS = {
	"custom_lr_gr_no",
	"custom_dispatch_from",
	"custom_assigned_to",
	# Transporter is seeded from the Pick List but editable in Draft; a change here
	# propagates back to the Pick List (delivery_note_on_update_draft).
	"custom_transporter_name",
	# vehicle_no is synced from the Pick List and read-only, so it's left out here.
	# Changes(HP) #43: Dispatch Date is editable in Draft; the save's on_update hook
	# (alpinos.dispatch_date_sync) carries it to the Sales Order, Pick List and Post Dispatch.
	"custom_dispatch_date",
}


_DN_QTY_EDIT_ROLES = {"Warehouse Admin", "Warehouse Manager", "System Manager", "PL Manager"}


def _can_edit_dn_qty():
	"""Only authorized roles may change DN item quantities; others re-post it unchanged."""
	return bool(set(frappe.get_roles()) & _DN_QTY_EDIT_ROLES)


def _apply_items_changes(dn, items):
	"""Apply qty edits and row removals from the page to dn.items."""
	if items is None:
		return

	items = json.loads(items) if isinstance(items, str) else items
	by_name = {row.name: row for row in dn.items}
	can_edit_qty = _can_edit_dn_qty()

	to_remove = []
	for entry in items:
		row_name = entry.get("name")
		if not row_name or row_name not in by_name:
			continue
		row = by_name[row_name]
		if entry.get("delete"):
			to_remove.append(row)
			continue
		# Qty edits from users without the role are ignored.
		if can_edit_qty and "qty" in entry and entry.get("qty") not in (None, ""):
			try:
				row.qty = float(entry["qty"])
			except (TypeError, ValueError):
				frappe.throw(f"Invalid quantity for row {row.idx}.")
		if "custom_remark" in entry and entry.get("custom_remark") is not None:
			row.custom_remark = (entry.get("custom_remark") or "").strip() or None

	for row in to_remove:
		dn.remove(row)


def _apply_dispatch_to_changes(dn, dispatch_to):
	"""Replace the Dispatch To child rows with the provided list."""
	if dispatch_to is None:
		return
	dispatch_to = json.loads(dispatch_to) if isinstance(dispatch_to, str) else dispatch_to

	dn.set("custom_dispatch_to", [])
	for entry in dispatch_to:
		text = (entry or {}).get("dispatch_to_address")
		if isinstance(text, str):
			text = text.strip()
		if not text:
			continue
		dn.append("custom_dispatch_to", {"dispatch_to_address": text})


@frappe.whitelist()
def save_delivery_note_data(name, header, items=None, dispatch_to=None):
	"""Save editable header fields, item edits and Dispatch To rows on a Draft DN."""
	header = json.loads(header) if isinstance(header, str) else header

	dn = frappe.get_doc("Delivery Note", name)
	dn.check_permission("write")

	if dn.docstatus != 0:
		frappe.throw("Submitted Delivery Note cannot be edited.")

	for k, v in header.items():
		if k not in _EDITABLE_HEADER_FIELDS:
			continue
		if k == "custom_dispatch_date":
			# A Datetime on the note, a date on the page: change the day, keep the time, and
			# never blank a mandatory date.
			if v:
				from alpinos.dispatch_date_sync import _stored

				dn.set(k, _stored("Delivery Note", getdate(v), dn.get(k)))
			continue
		dn.set(k, v if v not in ("", None) else None)

	_apply_items_changes(dn, items)
	_apply_dispatch_to_changes(dn, dispatch_to)

	dn.flags.ignore_mandatory = True
	dn.save(ignore_permissions=True)
	frappe.db.commit()
	return True


def _backfill_item_dates_from_pick_list(dn):
	"""Fill MFG/Expiry/Box/Batch on DN items from the DN row, else the Pick List Item, else the Batch master.

	Those DN date fields are reqd + read-only, so without this a DN can't submit
	when the Pick List left dates blank or predates them being copied.
	"""
	pl_row_names = [it.get("pick_list_item") for it in dn.items if it.get("pick_list_item")]
	pl_data = {}
	if pl_row_names:
		for r in frappe.get_all(
			"Pick List Item",
			filters={"name": ["in", pl_row_names]},
			fields=["name", "custom_mfg_date", "custom_expiry_date", "custom_box", "custom_batch_code"],
		):
			pl_data[r["name"]] = r

	changed = False

	def _fill(item, attr, value):
		nonlocal changed
		if value and not item.get(attr):
			item.set(attr, value)
			changed = True

	from alpinos.pick_list_api import _ensure_batch_exists

	for item in dn.items:
		pl = pl_data.get(item.get("pick_list_item")) or {}

		# Batch first: the free-text code lands in custom_batch_code; batch_no (a
		# Link to Batch) is set only when a real Batch exists, since a bare string
		# there fails DN submit.
		_fill(item, "custom_batch_code", pl.get("custom_batch_code"))
		if not item.get("batch_no") and item.get("custom_batch_code"):
			bn = _ensure_batch_exists(
				item.get("item_code"),
				item.custom_batch_code,
				pl.get("custom_mfg_date"),
				pl.get("custom_expiry_date"),
			)
			_fill(item, "batch_no", bn)
		_fill(item, "custom_box", pl.get("custom_box"))
		_fill(item, "custom_mfg_date", pl.get("custom_mfg_date"))
		_fill(item, "custom_expiry_date", pl.get("custom_expiry_date"))

		# Final fallback: read manufacturing / expiry from the Batch master.
		if item.get("batch_no") and (
			not item.get("custom_mfg_date") or not item.get("custom_expiry_date")
		):
			b = frappe.db.get_value(
				"Batch",
				item.batch_no,
				["manufacturing_date", "expiry_date"],
				as_dict=True,
			) or {}
			_fill(item, "custom_mfg_date", b.get("manufacturing_date"))
			_fill(item, "custom_expiry_date", b.get("expiry_date"))

	return changed


@frappe.whitelist()
def submit_delivery_note(name, header=None, items=None, dispatch_to=None):
	"""Save then submit the Delivery Note."""
	if header is not None:
		save_delivery_note_data(name, header, items, dispatch_to)

	dn = frappe.get_doc("Delivery Note", name)
	dn.check_permission("submit")
	if dn.docstatus == 0:
		if _backfill_item_dates_from_pick_list(dn):
			dn.flags.ignore_mandatory = True
			dn.save(ignore_permissions=True)
		dn.submit()
		frappe.db.commit()
	return dn.name


@frappe.whitelist()
def get_delivery_note_list(
	start=0,
	page_length=20,
	search="",
	status="",
	company="",
	sales_order="",
):
	if not frappe.has_permission("Delivery Note", "read"):
		frappe.throw(frappe._("You are not permitted to view Delivery Notes."), frappe.PermissionError)
	start = cint(start)
	page_length = cint(page_length)

	filters = {}
	if status:
		filters["status"] = status
	if company:
		filters["company"] = company
	if sales_order:
		filters["custom_sales_order_id"] = sales_order

	# A dedicated DN User only sees DNs assigned to them; admins/managers see all.
	_roles = set(frappe.get_roles())
	_override = {"System Manager", "Administrator", "Warehouse Admin", "Warehouse Manager", "PL Manager"}
	if "DN User" in _roles and not (_roles & _override):
		filters["custom_assigned_to"] = frappe.session.user

	# Changes(HP) #22: only notes of orders in the user's channels (frappe.get_all skips the
	# permission hooks that do this for the desk list).
	from alpinos.channel_access import allowed_sales_orders

	_allowed_sos = allowed_sales_orders()
	if _allowed_sos is not None:
		if sales_order:
			_allowed_sos = [sales_order] if sales_order in _allowed_sos else []
		filters["custom_sales_order_id"] = ["in", _allowed_sos or ["__no_match__"]]

	or_filters = []
	if search:
		or_filters = [
			["name", "like", f"%{search}%"],
			["custom_dn_so_customer_name", "like", f"%{search}%"],
			["customer_name", "like", f"%{search}%"],
		]

	rows = frappe.get_all(
		"Delivery Note",
		filters=filters,
		or_filters=or_filters,
		fields=[
			"name",
			"customer_name",
			"custom_dn_so_customer_name",
			"posting_date",
			"custom_dispatch_date",
			"company",
			"status",
			"docstatus",
			"custom_sales_order_id",
			"custom_transporter_name",
			"custom_lr_gr_no",
			"custom_invoice_no",
			"custom_assigned_to",
			"custom_total_boxes",
		],
		order_by="creation desc",
		limit_start=start,
		limit_page_length=page_length + 1,
	)

	has_more = len(rows) > page_length
	if has_more:
		rows = rows[:page_length]

	# Invoice No lives on the Sales Order (set after the DN is made); the DN's own
	# copy is usually empty, so show the SO's live value.
	so_ids = list({r.custom_sales_order_id for r in rows if r.get("custom_sales_order_id")})
	inv_by_so = {}
	if so_ids:
		for so in frappe.get_all(
			"Sales Order", filters={"name": ["in", so_ids]}, fields=["name", "custom_invoice_no"]
		):
			inv_by_so[so.name] = so.custom_invoice_no or ""

	for r in rows:
		r["custom_invoice_no"] = inv_by_so.get(r.get("custom_sales_order_id"), "")
		if r.get("custom_dispatch_date"):
			try:
				r["custom_dispatch_date"] = formatdate(str(r["custom_dispatch_date"])[:10])
			except Exception:
				pass
		if r.get("posting_date"):
			try:
				r["posting_date"] = formatdate(str(r["posting_date"]))
			except Exception:
				pass

	return {
		"data": rows,
		"has_more": has_more,
		"start": start,
		"page_length": page_length,
	}


# Bulk LR No. update (Warehouse Admin / Manager)
_LR_BULK_ROLES = {"Warehouse Admin", "Warehouse Manager", "System Manager"}


def _require_lr_roles():
	if not (set(frappe.get_roles()) & _LR_BULK_ROLES):
		frappe.throw(frappe._("Only Warehouse Admin / Manager can bulk-update LR No."))


def _so_po_invoice(so_id):
	"""(Customer PO, Invoice No) for a Sales Order — PO prefers the e-com PO number."""
	if not so_id:
		return "", ""
	r = frappe.db.get_value(
		"Sales Order", so_id, ["po_no", "custom_po_number", "custom_invoice_no"], as_dict=True
	) or {}
	return (r.get("custom_po_number") or r.get("po_no") or ""), (r.get("custom_invoice_no") or "")


# Changes(HP) #45: the bulk LR sheet carries the dispatch line as the Delivery Note list
# shows it. Only DISPATCH DATE and LR NO. are the warehouse's to fill — they are the two
# highlighted in yellow, and the only two read back on upload; every other column is the
# system's own figure, so editing one in the sheet changes nothing here.
_LR_COLUMNS = (
	("DATE", "posting_date"),
	("SO NO.", "custom_sales_order_id"),
	("CUSTOMER", "custom_dn_so_customer_name"),
	("PICKLIST PO NO.", "vehicle_no"),
	("PICK NO.", "pick_list"),
	("DISPATCH DATE", "custom_dispatch_date"),
	("TRANSPORTER", "custom_transporter_name"),
	("INVOICE ID", "invoice_id"),
	("LR NO.", "custom_lr_gr_no"),
	("Total Units", "custom_total_units_dn"),
	("Total Boxes", "custom_total_boxes"),
	("Gross Weight", "custom_dn_order_gross_weight"),
)
_LR_EDITABLE_COLUMNS = ("DISPATCH DATE", "LR NO.")
_LR_DATE_COLUMNS = ("DATE", "DISPATCH DATE")
_LR_COLUMN_WIDTHS = {
	"DATE": 12, "SO NO.": 18, "CUSTOMER": 30, "PICKLIST PO NO.": 18, "PICK NO.": 22,
	"DISPATCH DATE": 15, "TRANSPORTER": 20, "INVOICE ID": 18, "LR NO.": 18,
	"Total Units": 12, "Total Boxes": 12, "Gross Weight": 14,
}
_LR_HIGHLIGHT = "FFFF00"
_LR_DATE_FORMAT = "dd-mm-yyyy"


def _lr_key(label):
	"""A header as a comparable key: 'SO NO.' and 'So No' are the same column."""
	return "".join(ch for ch in str(label or "").upper() if ch.isalnum())


def _lr_pick_lists(dn_names):
	"""Pick List per Delivery Note ('PICK NO.'), joined when a note came from more than one."""
	if not dn_names:
		return {}
	rows = frappe.get_all(
		"Delivery Note Item",
		filters={"parent": ["in", dn_names], "against_pick_list": ["is", "set"]},
		fields=["parent", "against_pick_list"],
	)
	picks = {}
	for row in rows:
		names = picks.setdefault(row.parent, [])
		if row.against_pick_list not in names:
			names.append(row.against_pick_list)
	return {dn: ", ".join(names) for dn, names in picks.items()}


def _lr_rows_for_today():
	"""Draft Delivery Notes dispatching today, as the sheet's rows."""
	today = frappe.utils.today()
	# custom_dispatch_date is a Datetime, so match a full-day range, not "= today".
	dns = frappe.get_all(
		"Delivery Note",
		filters={
			"docstatus": 0,
			"custom_dispatch_date": ["between", [f"{today} 00:00:00", f"{today} 23:59:59"]],
		},
		fields=[
			"name", "posting_date", "custom_sales_order_id", "custom_dn_so_customer_name",
			"vehicle_no", "custom_dispatch_date", "custom_transporter_name", "custom_lr_gr_no",
			"custom_total_units_dn", "custom_total_boxes", "custom_dn_order_gross_weight",
		],
		order_by="name",
	)
	picks = _lr_pick_lists([dn.name for dn in dns])
	for dn in dns:
		dn["pick_list"] = picks.get(dn.name, "")
		dn["invoice_id"] = _so_po_invoice(dn.get("custom_sales_order_id"))[1]
		dn["posting_date"] = getdate(dn.get("posting_date")) if dn.get("posting_date") else None
		dn["custom_dispatch_date"] = (
			getdate(dn.get("custom_dispatch_date")) if dn.get("custom_dispatch_date") else None
		)
	return dns


@frappe.whitelist()
def download_lr_excel():
	"""Excel of DRAFT Delivery Notes dispatching TODAY, for bulk Dispatch Date / LR No. entry.

	Twelve columns in the order of the agreed format; DISPATCH DATE and LR NO. are the
	editable ones and are highlighted in yellow, the rest are filled in by the system.
	"""
	_require_lr_roles()
	import io

	try:
		import openpyxl
		from openpyxl.styles import Alignment, Font, PatternFill
	except Exception:
		frappe.throw(frappe._("openpyxl is required to build the LR Excel."))

	highlight = PatternFill("solid", fgColor=_LR_HIGHLIGHT)
	wb = openpyxl.Workbook()
	ws = wb.active
	ws.title = "LR Update"

	for col, (label, _field) in enumerate(_LR_COLUMNS, start=1):
		cell = ws.cell(row=1, column=col, value=label)
		cell.font = Font(bold=True)
		cell.alignment = Alignment(horizontal="center")
		if label in _LR_EDITABLE_COLUMNS:
			cell.fill = highlight
		ws.column_dimensions[cell.column_letter].width = _LR_COLUMN_WIDTHS.get(label, 16)

	for idx, dn in enumerate(_lr_rows_for_today(), start=2):
		for col, (label, field) in enumerate(_LR_COLUMNS, start=1):
			value = dn.get(field)
			cell = ws.cell(row=idx, column=col, value=value if value not in (None, "") else None)
			if label in _LR_DATE_COLUMNS:
				cell.number_format = _LR_DATE_FORMAT
			if label in _LR_EDITABLE_COLUMNS:
				cell.fill = highlight

	ws.freeze_panes = "A2"
	stream = io.BytesIO()
	wb.save(stream)

	frappe.response["filename"] = f"LR_Update_{frappe.utils.today()}.xlsx"
	frappe.response["filecontent"] = stream.getvalue()
	frappe.response["type"] = "binary"


def _lr_upload_date(value):
	"""A date out of an uploaded cell: a real Excel date, or a date someone typed."""
	import datetime

	if isinstance(value, datetime.datetime):
		return value.date()
	if isinstance(value, datetime.date):
		return value
	text = str(value or "").strip()
	if not text:
		return None
	text = text.split(" ")[0]
	for fmt in ("%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y", "%Y/%m/%d", "%d-%m-%y", "%d/%m/%y", "%m/%d/%Y"):
		try:
			return datetime.datetime.strptime(text, fmt).date()
		except ValueError:
			continue
	raise ValueError(text)


def _lr_upload_columns(header_row):
	"""Where SO NO., DISPATCH DATE and LR NO. sit in the uploaded sheet.

	Found by heading, so a sheet downloaded before #45 (Sales Order ID ... LR No.) still
	uploads; if the headings say nothing, fall back to that older layout's positions.
	"""
	columns = {"sales_order": None, "dispatch_date": None, "lr_no": None}
	for idx, label in enumerate(header_row or []):
		key = _lr_key(label)
		if key in ("SONO", "SALESORDERID") and columns["sales_order"] is None:
			columns["sales_order"] = idx
		elif key == "DISPATCHDATE" and columns["dispatch_date"] is None:
			columns["dispatch_date"] = idx
		elif key == "LRNO" and columns["lr_no"] is None:
			columns["lr_no"] = idx
	if columns["sales_order"] is None:
		columns["sales_order"] = 0
	if columns["lr_no"] is None:
		columns["lr_no"] = 3
	return columns


@frappe.whitelist()
def upload_lr_excel(file_url):
	"""Read a filled LR Excel, set Dispatch Date / LR No. on the matching draft DN by Sales
	Order ID, then submit it. Returns a summary and per-row failures.

	Only the two editable columns are read back: what the sheet says in the system's own
	columns is ignored, so an edit there changes nothing.
	"""
	_require_lr_roles()
	import io

	try:
		import openpyxl
	except Exception:
		frappe.throw(frappe._("openpyxl is required to read the uploaded Excel."))

	from frappe.utils.file_manager import get_file

	content = get_file(file_url)[1]
	wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
	ws = wb.active

	header = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), None)
	columns = _lr_upload_columns(header)

	def _cell(row, idx):
		if idx is None or not row or len(row) <= idx or row[idx] in (None, ""):
			return None
		return row[idx]

	updated, failed = 0, []
	for idx, r in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
		so_value = _cell(r, columns["sales_order"])
		so_id = str(so_value).strip() if so_value is not None else ""
		lr_value = _cell(r, columns["lr_no"])
		lr = str(lr_value).strip() if lr_value is not None else ""
		if not so_id:
			continue
		if not lr:
			failed.append({"row": idx, "sales_order": so_id, "reason": "LR No. is blank"})
			continue
		dispatch_value = _cell(r, columns["dispatch_date"])
		try:
			dispatch_date = _lr_upload_date(dispatch_value)
		except ValueError:
			failed.append({
				"row": idx, "sales_order": so_id,
				"reason": f"Dispatch Date '{dispatch_value}' is not a date",
			})
			continue
		dns = frappe.get_all(
			"Delivery Note", filters={"docstatus": 0, "custom_sales_order_id": so_id}, pluck="name"
		)
		if not dns:
			failed.append({"row": idx, "sales_order": so_id, "reason": "No draft Delivery Note found"})
			continue
		if len(dns) > 1:
			failed.append({"row": idx, "sales_order": so_id, "reason": "Multiple draft Delivery Notes — update individually"})
			continue
		try:
			dn = frappe.get_doc("Delivery Note", dns[0])
			if dispatch_date:
				# A Datetime on the note, a date in the sheet: change the day, keep the time.
				from alpinos.dispatch_date_sync import _stored

				dn.custom_dispatch_date = _stored("Delivery Note", dispatch_date, dn.custom_dispatch_date)
			dn.custom_lr_gr_no = lr
			dn.flags.ignore_permissions = True
			dn.submit()  # runs validate (LR mandatory now satisfied) + on_submit
			frappe.db.commit()
			updated += 1
		except Exception as e:
			frappe.db.rollback()
			failed.append({"row": idx, "sales_order": so_id, "reason": str(e)[:200]})

	return {
		"updated": updated,
		"failed": failed,
		"message": frappe._("{0} Delivery Notes updated and submitted successfully. {1} rows failed.").format(
			updated, len(failed)
		),
	}
