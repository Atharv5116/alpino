"""Take the charge for free rows off Delivery Notes made before they stopped being priced.

Until 76685a8, Marketing Freebies / Scheme / Additional Units rows reached a Delivery Note
with rate 0 but no price_list_rate, so ERPNext priced them from the Item Price list: the
note charged list price plus GST for samples the order never bills, and its grand total
ran above the order's (SOR-2627-02828: order 47,416.80, note 51,327.36). New notes carry
price_list_rate 0; this corrects the notes already made.

Each note is corrected on its OWN lines: the free rows go to 0 and the totals are
re-derived exactly as a save derives them (ERPNext's calculation, then the Sales Order
alignment where the branch has it). A part delivery therefore keeps its own share --
nothing is set to the order's total. A note is only rewritten when that same
calculation, run with the free rows untouched, reproduces its stored totals, so the free
rows are the only thing that can move; a note that does not is listed for review and
left alone.

What copied the old total is restated too: the Post Dispatch rows raised from the note,
and the order's Total Invoice Value -- on every picked or dispatched order, with the
order-rate rule of 76685a8.

Run by hand, dry run first:

  bench --site SITE execute alpinos.dn_free_row_repricing.preview   (writes nothing; prints
                                                                     and saves a workbook)
  bench --site SITE execute alpinos.dn_free_row_repricing.apply     (writes)
"""

import frappe
from frappe.utils import flt

FREE_TABLES = (
	"Sales Order Marketing Freebie",
	"Sales Order Scheme Item",
	"Sales Order Additional Units Item",
)

# Every price field a free row can carry on a Delivery Note Item.
_PRICE_FIELDS = (
	"price_list_rate", "base_price_list_rate",
	"margin_rate_or_amount", "rate_with_margin", "base_rate_with_margin",
	"discount_percentage", "discount_amount", "distributed_discount_amount",
	"rate", "base_rate", "net_rate", "base_net_rate", "stock_uom_rate",
	"amount", "base_amount", "net_amount", "base_net_amount",
)

_TOTALS = (
	"total", "net_total", "total_taxes_and_charges", "discount_amount",
	"grand_total", "rounded_total",
)

# A row is free when it was mapped from one of the free tables, or when the Pick List row
# it was picked on says it came from one (a row picked outside the order's Items table).
_FREE_ROW_SQL = """
	(
		EXISTS (SELECT 1 FROM `tabSales Order Marketing Freebie` f WHERE f.name = dni.so_detail)
		OR EXISTS (SELECT 1 FROM `tabSales Order Scheme Item` s WHERE s.name = dni.so_detail)
		OR EXISTS (SELECT 1 FROM `tabSales Order Additional Units Item` a WHERE a.name = dni.so_detail)
		OR EXISTS (SELECT 1 FROM `tabPick List Item` pli WHERE pli.name = dni.pick_list_item
		           AND IFNULL(pli.custom_source_table, 'Items') NOT IN ('', 'Items'))
	)
"""


def affected_notes():
	"""Non-return Delivery Notes (draft or submitted) with a free row that carries a price."""
	return frappe.db.sql(
		f"""
		SELECT DISTINCT dn.name
		FROM `tabDelivery Note Item` dni
		JOIN `tabDelivery Note` dn ON dn.name = dni.parent
		WHERE dn.docstatus < 2 AND IFNULL(dn.is_return, 0) = 0
		  AND (IFNULL(dni.amount, 0) <> 0 OR IFNULL(dni.rate, 0) <> 0
		       OR IFNULL(dni.price_list_rate, 0) <> 0)
		  AND {_FREE_ROW_SQL}
		ORDER BY dn.name
		""",
		pluck="name",
	)


def _free_row_names(delivery_note):
	return set(
		frappe.db.sql(
			f"""
			SELECT dni.name FROM `tabDelivery Note Item` dni
			WHERE dni.parent = %s AND {_FREE_ROW_SQL}
			""",
			delivery_note,
			pluck="name",
		)
	)


def _totals(doc):
	return {f: flt(doc.get(f), 2) for f in _TOTALS}


def _recalculate(doc):
	"""The value part of a Delivery Note save, without the rest of validate."""
	from alpinos import delivery_note_hooks

	doc.calculate_taxes_and_totals()
	# The Sales Order alignment is part of the save only on a branch that has it; where it
	# is absent the notes were saved without it, and recomputing without it matches them.
	align = getattr(delivery_note_hooks, "_align_value_with_sales_order", None)
	if align:
		align(doc)


def _sales_orders_of(doc):
	names = {(doc.get("custom_sales_order_id") or "").strip()}
	names.update(row.against_sales_order for row in doc.items if row.get("against_sales_order"))
	names.discard("")
	names.discard(None)
	return names


def reprice(delivery_note, apply=False):
	"""Correct one note. Returns what happened to it (and writes only when apply)."""
	doc = frappe.get_doc("Delivery Note", delivery_note)
	stored = _totals(doc)
	free = _free_row_names(delivery_note)
	rows = [r for r in doc.items if r.name in free]
	if not rows:
		return {"delivery_note": delivery_note, "status": "no free rows"}
	orders = sorted(_sales_orders_of(doc))
	primary = (doc.get("custom_sales_order_id") or "").strip() or (orders[0] if orders else "")

	# Guard: the same calculation over the untouched note must give back what is stored.
	# If it does not, the note was changed some other way, and rewriting it would move
	# more than the free rows.
	_recalculate(doc)
	recomputed = _totals(doc)
	drift = {f: (stored[f], recomputed[f]) for f in _TOTALS if abs(stored[f] - recomputed[f]) > 0.005}
	if drift:
		return {"delivery_note": delivery_note, "status": "review", "drift": drift,
		        "docstatus": doc.docstatus, "sales_order": primary, "sales_orders": orders,
		        "old_grand_total": stored["grand_total"]}

	charged = flt(sum(flt(r.amount) for r in rows), 2)
	for row in rows:
		for field in _PRICE_FIELDS:
			if row.meta.has_field(field):
				row.set(field, 0)
	_recalculate(doc)
	if hasattr(doc, "set_total_in_words"):
		doc.set_total_in_words()
	after = _totals(doc)

	result = {
		"delivery_note": delivery_note,
		"status": "corrected" if apply else "would correct",
		"docstatus": doc.docstatus,
		"free_rows": len(rows),
		"free_rows_charged": charged,
		"old_grand_total": stored["grand_total"],
		"new_grand_total": after["grand_total"],
		"sales_order": primary,
		"sales_orders": orders,
	}
	if not apply:
		return result

	# Selling values only: a Delivery Note posts stock at valuation, so no ledger moves.
	doc.db_update()
	for row in list(doc.items) + list(doc.get("taxes") or []):
		row.db_update()
	doc.add_comment(
		"Info",
		"Free rows (Marketing Freebies / Scheme / Additional Units) had been priced from the "
		"Item Price list; {0} row(s), {1} before GST, set to 0. Grand total {2} -> {3}.".format(
			len(rows),
			frappe.format_value(charged, {"fieldtype": "Currency"}),
			frappe.format_value(stored["grand_total"], {"fieldtype": "Currency"}),
			frappe.format_value(after["grand_total"], {"fieldtype": "Currency"}),
		),
	)
	return result


def run(apply=0):
	"""Correct every affected note, then restate Post Dispatch and the orders' value."""
	from alpinos.so_invoice_value import (
		refresh_for_sales_order,
		refresh_post_dispatch_for_delivery_note,
	)

	apply = bool(int(apply))
	results = [reprice(name, apply=apply) for name in affected_notes()]
	fixed = [r for r in results if r["status"] in ("corrected", "would correct")]
	review = [r for r in results if r["status"] == "review"]

	orders = sorted({so for r in fixed for so in r["sales_orders"]})
	post_dispatch_rows = 0
	if apply:
		for r in fixed:
			post_dispatch_rows += refresh_post_dispatch_for_delivery_note(r["delivery_note"])
		for so in orders:
			refresh_for_sales_order(so)
		frappe.db.commit()

	return {
		"mode": "APPLY" if apply else "DRY-RUN",
		"notes_with_priced_free_rows": len(results),
		"corrected": len(fixed),
		"overcharge_removed_before_gst": flt(sum(r["free_rows_charged"] for r in fixed), 2),
		"grand_total_removed": flt(sum(r["old_grand_total"] - r["new_grand_total"] for r in fixed), 2),
		"sales_orders_restated": len(orders),
		"post_dispatch_rows_restated": post_dispatch_rows,
		"left_for_review": review,
		"notes": fixed[:50],
	}


# --------------------------------------------------------------- dry run for review

_NOTE_HEADERS = (
	"Sales Order", "Customer", "Order Amount", "Delivery Note", "Status", "Units (note / order)",
	"Note Amount Now", "Note Amount After", "Reduced By", "Free Rows",
)
_ORDER_HEADERS = (
	"Sales Order", "Customer", "Order Amount", "Delivery", "Submitted Notes Now",
	"Submitted Notes After", "Total Invoice Value Now", "Total Invoice Value After",
)
_REVIEW_HEADERS = (
	"Sales Order", "Customer", "Order Amount", "Delivery Note", "Status", "Note Amount", "Why",
)
_MONEY = {
	"Order Amount", "Note Amount Now", "Note Amount After", "Reduced By", "Submitted Notes Now",
	"Submitted Notes After", "Total Invoice Value Now", "Total Invoice Value After", "Note Amount",
}


def _order_facts(sales_orders):
	"""Per order: its amount, customer, units ordered and delivered, the submitted notes'
	current total, and the stored Total Invoice Value."""
	if not sales_orders:
		return {}
	rows = frappe.db.sql(
		"""
		SELECT so.name, so.customer_name, so.grand_total,
		       IFNULL(so.custom_total_invoice_value, 0) AS tiv,
		       (SELECT IFNULL(SUM(i.stock_qty), 0) FROM `tabSales Order Item` i
		        WHERE i.parent = so.name) AS ordered,
		       (SELECT IFNULL(SUM(dni.stock_qty), 0)
		        FROM `tabDelivery Note Item` dni
		        JOIN `tabDelivery Note` dn ON dn.name = dni.parent
		             AND dn.docstatus = 1 AND IFNULL(dn.is_return, 0) = 0
		        JOIN `tabSales Order Item` i ON i.name = dni.so_detail AND i.parent = so.name) AS delivered,
		       (SELECT IFNULL(SUM(dn.grand_total), 0) FROM `tabDelivery Note` dn
		        WHERE dn.custom_sales_order_id = so.name
		          AND dn.docstatus = 1 AND IFNULL(dn.is_return, 0) = 0) AS notes_total
		FROM `tabSales Order` so
		WHERE so.name IN %(names)s
		""",
		{"names": tuple(sales_orders)},
		as_dict=True,
	)
	return {r.name: r for r in rows}


def _note_units(delivery_notes):
	"""Units each note carries on the order's own lines (free rows not counted)."""
	if not delivery_notes:
		return {}
	return dict(
		frappe.db.sql(
			"""
			SELECT dni.parent, SUM(dni.stock_qty)
			FROM `tabDelivery Note Item` dni
			JOIN `tabSales Order Item` i ON i.name = dni.so_detail
			WHERE dni.parent IN %(names)s
			GROUP BY dni.parent
			""",
			{"names": tuple(delivery_notes)},
		)
	)


def _qty(value):
	return f"{flt(value):g}"


def _print_table(title, headers, rows, limit=100):
	print(f"\n{title}  ({len(rows)})")
	if not rows:
		print("  none")
		return
	shown = rows[:limit]
	cells = [[f"{v:,.2f}" if h in _MONEY else str(v if v is not None else "") for h, v in zip(headers, r)]
	         for r in shown]
	widths = [max(len(h), *(len(c[i]) for c in cells)) for i, h in enumerate(headers)]
	line = lambda vals: "  " + "  ".join(
		v.rjust(w) if h in _MONEY else v.ljust(w) for h, v, w in zip(headers, vals, widths))
	print(line(headers))
	print("  " + "  ".join("-" * w for w in widths))
	for c in cells:
		print(line(c))
	if len(rows) > limit:
		print(f"  ... and {len(rows) - limit} more in the workbook")


def _save_workbook(path, sheets):
	from openpyxl import Workbook
	from openpyxl.styles import Font

	wb = Workbook()
	wb.remove(wb.active)
	for title, headers, rows in sheets:
		ws = wb.create_sheet(title[:31])
		ws.append(list(headers))
		for cell in ws[1]:
			cell.font = Font(bold=True)
		for r in rows:
			ws.append(list(r))
		for i, h in enumerate(headers, start=1):
			letter = ws.cell(row=1, column=i).column_letter
			width = max([len(h)] + [len(f"{r[i - 1]:,.2f}" if h in _MONEY else str(r[i - 1] or "")) for r in rows])
			ws.column_dimensions[letter].width = min(width + 2, 60)
			if h in _MONEY:
				for row in ws.iter_rows(min_row=2, min_col=i, max_col=i):
					row[0].number_format = "#,##0.00"
		ws.freeze_panes = "A2"
	wb.save(path)


def preview(file_path=None):
	"""Dry run for review. Writes nothing to the database.

	Lists every Delivery Note that would be corrected, every order whose Total Invoice
	Value would be restated, and any note left for review -- each beside the Sales Order's
	own amount -- prints them and saves them as a workbook (default: the site's
	private/files/delivery_note_amount_dry_run.xlsx).
	"""
	from alpinos.so_invoice_value import backfill

	results = [reprice(name, apply=False) for name in affected_notes()]
	fixed = [r for r in results if r["status"] == "would correct"]
	review = [r for r in results if r["status"] == "review"]
	restated = {r["sales_order"]: r for r in backfill(apply=0, sample=0)["sample"]}

	orders = {r["sales_order"] for r in fixed + review if r.get("sales_order")} | set(restated)
	facts = _order_facts(orders)
	units = _note_units([r["delivery_note"] for r in fixed])
	blank = frappe._dict(customer_name="", grand_total=0, tiv=0, ordered=0, delivered=0, notes_total=0)

	# Only submitted notes count towards the order's dispatched total.
	reduction = {}
	for r in fixed:
		if r["docstatus"] == 1:
			reduction[r["sales_order"]] = reduction.get(r["sales_order"], 0) + r["old_grand_total"] - r["new_grand_total"]

	note_rows = []
	for r in sorted(fixed, key=lambda x: (x["sales_order"], x["delivery_note"])):
		f = facts.get(r["sales_order"], blank)
		note_rows.append((
			r["sales_order"], f.customer_name, flt(f.grand_total, 2), r["delivery_note"],
			"Submitted" if r["docstatus"] == 1 else "Draft",
			f"{_qty(units.get(r['delivery_note']))} / {_qty(f.ordered)}",
			r["old_grand_total"], r["new_grand_total"],
			flt(r["old_grand_total"] - r["new_grand_total"], 2), r["free_rows"],
		))

	order_rows = []
	for so in sorted(orders):
		f = facts.get(so, blank)
		tiv_now = flt(f.tiv, 2)
		tiv_after = flt(restated[so]["new"], 2) if so in restated else tiv_now
		delivery = ("Full" if flt(f.delivered) >= flt(f.ordered) and flt(f.ordered)
		            else f"Part ({_qty(f.delivered)} of {_qty(f.ordered)})")
		order_rows.append((
			so, f.customer_name, flt(f.grand_total, 2), delivery, flt(f.notes_total, 2),
			flt(flt(f.notes_total) - reduction.get(so, 0), 2), tiv_now, tiv_after,
		))

	review_rows = []
	for r in sorted(review, key=lambda x: (x["sales_order"], x["delivery_note"])):
		f = facts.get(r["sales_order"], blank)
		review_rows.append((
			r["sales_order"], f.customer_name, flt(f.grand_total, 2), r["delivery_note"],
			"Submitted" if r["docstatus"] == 1 else "Draft", r["old_grand_total"],
			"saved totals do not match a recalculation; check by hand",
		))

	import os

	path = os.path.abspath(
		file_path or frappe.get_site_path("private", "files", "delivery_note_amount_dry_run.xlsx")
	)
	_save_workbook(path, [
		("Delivery Notes", _NOTE_HEADERS, note_rows),
		("Sales Orders", _ORDER_HEADERS, order_rows),
		("Left for review", _REVIEW_HEADERS, review_rows),
	])

	summary = {
		"mode": "DRY-RUN (nothing written)",
		"delivery_notes_to_correct": len(fixed),
		"of_which_draft": sum(1 for r in fixed if r["docstatus"] == 0),
		"amount_removed_from_notes": flt(sum(r["old_grand_total"] - r["new_grand_total"] for r in fixed), 2),
		"sales_orders_whose_total_invoice_value_changes": sum(1 for r in order_rows if r[6] != r[7]),
		"left_for_review": len(review),
		"workbook": path,
	}
	print("\nDRY RUN -- nothing has been written.")
	_print_table("Delivery Notes that charged for free rows", _NOTE_HEADERS, note_rows)
	_print_table("Sales Orders: dispatched total and Total Invoice Value", _ORDER_HEADERS, order_rows)
	_print_table("Left for review (will NOT be changed)", _REVIEW_HEADERS, review_rows)
	print(f"\nWorkbook: {path}\n")
	return summary


def apply():
	"""Write what preview() shows: correct the notes, restate their Post Dispatch rows, and
	restate Total Invoice Value on every picked or dispatched order."""
	from alpinos.so_invoice_value import backfill

	notes = run(apply=1)
	orders = backfill(apply=1, sample=0)
	return {
		"mode": "APPLIED",
		"delivery_notes_corrected": notes["corrected"],
		"amount_removed_from_notes": notes["grand_total_removed"],
		"post_dispatch_rows_restated": notes["post_dispatch_rows_restated"],
		"sales_orders_whose_total_invoice_value_changed": orders["changed"],
		"left_for_review": [r["delivery_note"] for r in notes["left_for_review"]],
	}
