"""Changes(HP) #45: the Delivery Note bulk LR Excel.

The sheet carries twelve columns in the agreed order. DISPATCH DATE and LR NO. are the
warehouse's to fill and are highlighted in yellow; the other ten are the system's own
figures and are ignored when the sheet comes back, so editing one of them changes
nothing. Everything else about the download (draft notes dispatching today) and the
upload (match the order, set the LR No., submit) stays as it was.

The download runs for real and the workbook it produces is read back with openpyxl; the
upload runs for real too, on a real draft Delivery Note of the site, with the uploaded
file handed to it in memory instead of through a File document. One transaction, rolled
back at the end, and commits made by the code under test are ignored while it runs.

Run:  bench --site alpinos.test execute alpinos.lr_excel_test.run
"""

import datetime
import io

import frappe
from frappe.utils import flt, getdate, today

from alpinos.alpinos_development.page.delivery_note_entry import delivery_note_entry as dne

R = []

REFERENCE_HEADERS = [
	"DATE", "SO NO.", "CUSTOMER", "PICKLIST PO NO.", "PICK NO.", "DISPATCH DATE",
	"TRANSPORTER", "INVOICE ID", "LR NO.", "Total Units", "Total Boxes", "Gross Weight",
]
YELLOW = "FFFF00"


def check(label, fn):
	"""Run one case and undo it: a case that submits a note must not leave it submitted
	for the next one, whether it passed or not."""
	try:
		fn()
		R.append(("PASS", label, ""))
	except AssertionError as e:
		R.append(("FAIL", label, str(e)))
	except Exception as e:
		R.append(("ERROR", label, f"{type(e).__name__}: {str(e)[:300]}"))
	finally:
		frappe.db.rollback()


def _assert(cond, msg=""):
	if not cond:
		raise AssertionError(msg)


def _sheet():
	"""Run the real download and read the workbook it put on the response."""
	import openpyxl

	frappe.response.pop("filecontent", None)
	dne.download_lr_excel()
	_assert(frappe.response.get("filecontent"), "download produced no file")
	wb = openpyxl.load_workbook(io.BytesIO(frappe.response["filecontent"]))
	return wb.active


def _fill(cell):
	"""The cell's background colour, '' when it has none."""
	if not cell.fill or cell.fill.fill_type in (None, "none"):
		return ""
	colour = cell.fill.start_color.rgb if cell.fill.start_color else None
	return (colour or "")[-6:]


def _row_for(ws, dn_name, so_id):
	for row in ws.iter_rows(min_row=2):
		if str(row[1].value or "") == so_id:
			return row
	raise AssertionError(f"{dn_name} ({so_id}) is not in the sheet")


def _upload(rows, headers=None):
	"""Feed the real upload a workbook built from `rows`, without touching the disk."""
	import openpyxl
	from frappe.utils import file_manager

	wb = openpyxl.Workbook()
	ws = wb.active
	ws.append(headers if headers is not None else REFERENCE_HEADERS)
	for row in rows:
		ws.append(row)
	stream = io.BytesIO()
	wb.save(stream)

	real_get_file = file_manager.get_file
	file_manager.get_file = lambda file_url: ("lr.xlsx", stream.getvalue())
	try:
		return dne.upload_lr_excel("/files/lr.xlsx")
	finally:
		file_manager.get_file = real_get_file


def _draft_dn(dispatch_on=None):
	"""A draft Delivery Note of the site, dispatching on the given day and able to submit.

	The upload submits the note, so the note needs the fields every Delivery Note needs
	and stock behind its rows: whatever this one is missing is filled in here, inside the
	transaction, so the test is about the LR sheet and not about a half-filled fixture.
	"""
	name = frappe.db.sql(
		"""
		SELECT dn.name FROM `tabDelivery Note` dn
		WHERE dn.docstatus = 0 AND IFNULL(dn.custom_sales_order_id, '') != ''
			AND (SELECT COUNT(*) FROM `tabDelivery Note` o
				 WHERE o.docstatus = 0 AND o.custom_sales_order_id = dn.custom_sales_order_id) = 1
		ORDER BY dn.creation DESC LIMIT 1
		"""
	)
	if not name:
		return None
	dn = frappe.get_doc("Delivery Note", name[0][0])
	frappe.db.set_single_value("Stock Settings", "allow_negative_stock", 1)
	header = {"custom_dispatch_date": f"{dispatch_on or today()} 09:30:00"}
	if not dn.custom_transporter_name:
		header["custom_transporter_name"] = "LR Test Transporter"
	if not dn.vehicle_no:
		header["vehicle_no"] = "LR-TEST-PO"
	if not dn.custom_dispatch_from:
		header["custom_dispatch_from"] = "LR Test Dispatch From"
	if not dn.custom_delivery_date:
		header["custom_delivery_date"] = today()
	frappe.db.set_value("Delivery Note", dn.name, header, update_modified=False)
	if not frappe.db.exists("Alpinos DN Dispatch To", {"parent": dn.name}):
		row = frappe.get_doc({
			"doctype": "Alpinos DN Dispatch To", "parent": dn.name, "parenttype": "Delivery Note",
			"parentfield": "custom_dispatch_to", "idx": 1, "dispatch_to_address": "LR Test Dispatch To",
		})
		row.db_insert()
	dn.reload()
	return dn


def run():
	R.clear()
	frappe.set_user("Administrator")
	real_commit = frappe.db.commit
	frappe.db.commit = lambda *args, **kwargs: None
	try:
		dn = _draft_dn()
		if not dn:
			R.append(("SKIP", "#45 LR Excel", "no draft Delivery Note with a Sales Order on this site"))
			return _report()
		so_id = dn.custom_sales_order_id

		def headers_match_the_reference():
			_draft_dn()
			ws = _sheet()
			headers = [ws.cell(row=1, column=c).value for c in range(1, len(REFERENCE_HEADERS) + 1)]
			_assert(headers == REFERENCE_HEADERS, f"headers {headers}")
			_assert(ws.max_column == len(REFERENCE_HEADERS),
				f"sheet has {ws.max_column} columns, the format has {len(REFERENCE_HEADERS)}")

		def only_the_editable_columns_are_yellow():
			dn = _draft_dn()
			ws = _sheet()
			row = _row_for(ws, dn.name, so_id)
			for idx, label in enumerate(REFERENCE_HEADERS):
				head = ws.cell(row=1, column=idx + 1)
				wanted = YELLOW if label in ("DISPATCH DATE", "LR NO.") else ""
				_assert(_fill(head) == wanted, f"header {label}: fill {_fill(head)!r}, wanted {wanted!r}")
				_assert(_fill(row[idx]) == wanted, f"cell {label}: fill {_fill(row[idx])!r}, wanted {wanted!r}")

		def system_columns_carry_the_system_values():
			dn = _draft_dn()
			ws = _sheet()
			row = _row_for(ws, dn.name, so_id)
			invoice = frappe.db.get_value("Sales Order", so_id, "custom_invoice_no") or ""
			picks = frappe.get_all(
				"Delivery Note Item", filters={"parent": dn.name, "against_pick_list": ["is", "set"]},
				pluck="against_pick_list", distinct=True)
			expected = {
				"DATE": getdate(dn.posting_date),
				"SO NO.": so_id,
				"CUSTOMER": dn.custom_dn_so_customer_name or None,
				"PICKLIST PO NO.": dn.vehicle_no or None,
				"PICK NO.": ", ".join(sorted(set(picks))) if picks else None,
				"DISPATCH DATE": getdate(dn.custom_dispatch_date),
				"TRANSPORTER": dn.custom_transporter_name or None,
				"INVOICE ID": invoice or None,
				"LR NO.": dn.custom_lr_gr_no or None,
				"Total Units": flt(dn.custom_total_units_dn),
				"Total Boxes": flt(dn.custom_total_boxes),
				"Gross Weight": flt(dn.custom_dn_order_gross_weight),
			}
			for idx, label in enumerate(REFERENCE_HEADERS):
				got = row[idx].value
				if isinstance(got, datetime.datetime):
					got = got.date()
				want = expected[label]
				if isinstance(want, (int, float)) and isinstance(got, (int, float)):
					_assert(flt(got) == flt(want), f"{label}: {got} != {want}")
				elif label == "PICK NO." and got:
					_assert(sorted(str(got).split(", ")) == sorted(str(want).split(", ")), f"{label}: {got} != {want}")
				else:
					_assert(got == want, f"{label}: {got!r} != {want!r}")

		def dates_are_real_dates():
			dn = _draft_dn()
			ws = _sheet()
			row = _row_for(ws, dn.name, so_id)
			for label in ("DATE", "DISPATCH DATE"):
				cell = row[REFERENCE_HEADERS.index(label)]
				_assert(isinstance(cell.value, (datetime.date, datetime.datetime)),
					f"{label} is {type(cell.value).__name__}, not a date Excel can edit")
				_assert(cell.number_format == "dd-mm-yyyy", f"{label} format {cell.number_format}")

		def only_notes_dispatching_today_are_listed():
			dn = _draft_dn()
			ws = _sheet()
			listed = {str(r[1].value or "") for r in ws.iter_rows(min_row=2)}
			_assert(so_id in listed, "today's draft note is missing from the sheet")
			frappe.db.set_value("Delivery Note", dn.name, "custom_dispatch_date",
				f"{frappe.utils.add_days(today(), 3)} 09:30:00", update_modified=False)
			try:
				later = {str(r[1].value or "") for r in _sheet().iter_rows(min_row=2)}
				_assert(so_id not in later, "a note dispatching in three days is still on today's sheet")
			finally:
				frappe.db.set_value("Delivery Note", dn.name, "custom_dispatch_date",
					f"{today()} 09:30:00", update_modified=False)

		def upload_sets_both_editable_columns():
			_draft_dn()
			day = frappe.utils.add_days(today(), 1)
			result = _upload([[today(), so_id, "EDITED NAME", "EDITED PO", "EDITED PICK", day,
				"EDITED TRANSPORTER", "EDITED INVOICE", "LR-45-A", 9999, 8888, 7777]])
			_assert(result["updated"] == 1, f"updated={result['updated']} failed={result['failed']}")
			after = frappe.get_doc("Delivery Note", dn.name)
			_assert(after.custom_lr_gr_no == "LR-45-A", f"LR No. is {after.custom_lr_gr_no!r}")
			_assert(getdate(after.custom_dispatch_date) == getdate(day),
				f"dispatch date is {after.custom_dispatch_date}, wanted {day}")
			_assert(str(after.custom_dispatch_date)[11:19] == "09:30:00",
				f"the note's dispatch time was lost: {after.custom_dispatch_date}")
			_assert(after.docstatus == 1, "the note was not submitted")

		def the_dispatch_date_reaches_the_order():
			_draft_dn()
			day = frappe.utils.add_days(today(), 2)
			result = _upload([[today(), so_id, "", "", "", day, "", "", "LR-45-B", "", "", ""]])
			_assert(result["updated"] == 1, f"updated={result['updated']} failed={result['failed']}")
			so_date = frappe.db.get_value("Sales Order", so_id, "custom_dispatch_date")
			_assert(so_date and getdate(so_date) == getdate(day),
				f"order's dispatch date is {so_date}, wanted {day}")

		def fixed_columns_do_not_update_the_system():
			before = _draft_dn()
			result = _upload([["2001-01-01", so_id, "EDITED NAME", "EDITED PO", "EDITED PICK",
				before.custom_dispatch_date, "EDITED TRANSPORTER", "EDITED INVOICE", "LR-45-C",
				9999, 8888, 7777]])
			_assert(result["updated"] == 1, f"updated={result['updated']} failed={result['failed']}")
			after = frappe.get_doc("Delivery Note", dn.name)
			_assert(after.custom_dn_so_customer_name == before.custom_dn_so_customer_name,
				f"customer changed to {after.custom_dn_so_customer_name!r}")
			_assert(after.vehicle_no == before.vehicle_no, f"Picklist PO No. changed to {after.vehicle_no!r}")
			_assert(after.custom_transporter_name == before.custom_transporter_name,
				f"transporter changed to {after.custom_transporter_name!r}")
			_assert(flt(after.custom_total_units_dn) == flt(before.custom_total_units_dn),
				"Total Units changed")
			_assert(flt(after.custom_total_boxes) == flt(before.custom_total_boxes), "Total Boxes changed")
			_assert(getdate(after.posting_date) != getdate("2001-01-01"),
				"the sheet's DATE reached the note's posting date")
			invoice = frappe.db.get_value("Sales Order", so_id, "custom_invoice_no") or ""
			_assert(invoice != "EDITED INVOICE", "the sheet's Invoice ID reached the order")

		def a_blank_dispatch_date_keeps_the_note_s_own():
			_draft_dn()
			before = frappe.db.get_value("Delivery Note", dn.name, "custom_dispatch_date")
			result = _upload([[today(), so_id, "", "", "", "", "", "", "LR-45-D", "", "", ""]])
			_assert(result["updated"] == 1, f"updated={result['updated']} failed={result['failed']}")
			after = frappe.db.get_value("Delivery Note", dn.name, "custom_dispatch_date")
			_assert(str(after) == str(before), f"dispatch date changed from {before} to {after}")

		def a_blank_lr_still_fails_the_row():
			_draft_dn()
			result = _upload([[today(), so_id, "", "", "", today(), "", "", "", "", "", ""]])
			_assert(result["updated"] == 0, "a row without an LR No. was accepted")
			_assert(result["failed"] and result["failed"][0]["reason"] == "LR No. is blank",
				f"failures: {result['failed']}")
			_assert(frappe.db.get_value("Delivery Note", dn.name, "docstatus") == 0,
				"the note was submitted although the LR No. was blank")

		def a_bad_dispatch_date_fails_the_row():
			_draft_dn()
			result = _upload([[today(), so_id, "", "", "", "not a date", "", "", "LR-45-E", "", "", ""]])
			_assert(result["updated"] == 0, "a row with an unreadable date was accepted")
			_assert(result["failed"] and "is not a date" in result["failed"][0]["reason"],
				f"failures: {result['failed']}")
			_assert(frappe.db.get_value("Delivery Note", dn.name, "docstatus") == 0,
				"the note was submitted although the dispatch date was unreadable")

		def a_sheet_from_before_45_still_uploads():
			_draft_dn()
			result = _upload(
				[[so_id, "PO-OLD", "INV-OLD", "LR-45-F"]],
				headers=["Sales Order ID", "Customer PO / PO Number", "Invoice No.", "LR No."],
			)
			_assert(result["updated"] == 1, f"updated={result['updated']} failed={result['failed']}")
			_assert(frappe.db.get_value("Delivery Note", dn.name, "custom_lr_gr_no") == "LR-45-F",
				"the older four-column sheet no longer sets the LR No.")

		def typed_dates_are_read_day_first():
			_assert(dne._lr_upload_date("15-09-2026") == datetime.date(2026, 9, 15), "15-09-2026")
			_assert(dne._lr_upload_date("2026-09-15") == datetime.date(2026, 9, 15), "2026-09-15")
			_assert(dne._lr_upload_date("05/09/2026") == datetime.date(2026, 9, 5), "05/09/2026")
			_assert(dne._lr_upload_date(datetime.datetime(2026, 9, 15, 9, 30)) == datetime.date(2026, 9, 15),
				"an Excel date")
			_assert(dne._lr_upload_date("") is None and dne._lr_upload_date(None) is None, "blank")

		check("#45 the sheet has the reference columns, in order", headers_match_the_reference)
		check("#45 only DISPATCH DATE and LR NO. are yellow", only_the_editable_columns_are_yellow)
		check("#45 the system columns carry the system's own values", system_columns_carry_the_system_values)
		check("#45 DATE and DISPATCH DATE are real Excel dates", dates_are_real_dates)
		check("#45 the sheet still lists today's draft notes only", only_notes_dispatching_today_are_listed)
		check("#45 upload sets Dispatch Date and LR No., then submits", upload_sets_both_editable_columns)
		check("#45 the uploaded dispatch date reaches the Sales Order", the_dispatch_date_reaches_the_order)
		check("#45 edits to the fixed columns update nothing", fixed_columns_do_not_update_the_system)
		check("#45 a blank dispatch date keeps the note's own", a_blank_dispatch_date_keeps_the_note_s_own)
		check("#45 a blank LR No. still fails the row", a_blank_lr_still_fails_the_row)
		check("#45 an unreadable dispatch date fails the row", a_bad_dispatch_date_fails_the_row)
		check("#45 a sheet downloaded before #45 still uploads", a_sheet_from_before_45_still_uploads)
		check("#45 typed dates are read day-first", typed_dates_are_read_day_first)
		return _report()
	finally:
		frappe.db.commit = real_commit
		frappe.db.rollback()
		frappe.set_user("Administrator")


def _report():
	width = max(len(r[1]) for r in R)
	for status, label, detail in R:
		print(f"[{status}] {label.ljust(width)}  {detail}")
	print(f"{sum(1 for r in R if r[0] == 'PASS')}/{sum(1 for r in R if r[0] != 'SKIP')} passed"
	      + (f", {sum(1 for r in R if r[0] == 'SKIP')} skipped" if any(r[0] == "SKIP" for r in R) else ""))
	return R
